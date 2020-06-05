import openpyxl.chart

wb = openpyxl.Workbook()
sheet = wb.active
for i in range(1, 11):
    sheet["A" + str(i)] = i

refObj = openpyxl.chart.reference.Reference(sheet, 1, 1, 10, 1)

seriesObj = openpyxl.chart.Series(refObj, title="First series")

chartObj = openpyxl.chart.BarChart()
chartObj.append(seriesObj)
sheet.add_chart(chartObj, "E5")
# chartObj.drawing.top = 50  # set the position
# chartObj.drawing.left = 100
# chartObj.drawing.width = 300  # set the size
# chartObj.drawing.height = 200

# sheet.add_charts(chartObj)

wb.save("sampleChart.xlsx")
